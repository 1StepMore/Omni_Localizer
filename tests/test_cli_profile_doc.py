"""Tests for the `ol profile-doc` CLI subcommand."""
from __future__ import annotations

import json
from pathlib import Path

import pytest
from typer.testing import CliRunner

from ol_cli import app

runner = CliRunner()


@pytest.fixture
def sample_md(tmp_path):
    """Create a sample markdown file."""
    md = tmp_path / "doc.md"
    md.write_text(
        "# Hello World\n\nThis is a test document about engineering best practices.\n",
        encoding="utf-8",
    )
    return md


class TestProfileDocCLI:
    """Test the ol profile-doc CLI command."""

    def test_cli_runs_with_minimum_args(self, sample_md):
        result = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
        ])
        assert result.exit_code == 0
        output = json.loads(result.stdout)
        assert "tone" in output
        assert "register" in output
        assert "summary" in output

    def test_cli_with_output_file(self, sample_md, tmp_path):
        output_path = tmp_path / "profile.json"
        result = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
            "--output", str(output_path),
        ])
        assert result.exit_code == 0
        assert output_path.exists()
        profile = json.loads(output_path.read_text(encoding="utf-8"))
        assert "tone" in profile

    def test_cli_with_source_lang(self, sample_md):
        result = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
            "--source-lang", "zh",
        ])
        assert result.exit_code == 0
        output = json.loads(result.stdout)
        # Standard fields are present in FAKE_LLM mode
        assert "tone" in output
        assert "register" in output
        assert "summary" in output

    def test_cli_with_cache_dir(self, sample_md, tmp_path):
        cache_dir = tmp_path / "cache"
        # First call: cache miss
        result1 = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
            "--cache-dir", str(cache_dir),
        ])
        assert result1.exit_code == 0
        # Second call: cache hit (should be fast)
        result2 = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
            "--cache-dir", str(cache_dir),
        ])
        assert result2.exit_code == 0
        # Both outputs should be identical
        out1 = json.loads(result1.output)
        out2 = json.loads(result2.output)
        assert out1 == out2
        # Cache directory should have at least one file
        assert any(cache_dir.glob("*.json"))

    def test_cli_missing_file(self, tmp_path):
        result = runner.invoke(app, [
            "profile-doc",
            str(tmp_path / "nonexistent.md"),
        ])
        assert result.exit_code != 0

    def test_cli_help(self):
        result = runner.invoke(app, ["profile-doc", "--help"])
        assert result.exit_code == 0
        assert "profile-doc" in result.output


class TestProfileDocCLIWithConfig:
    """Test --config option (loads profiling role from config)."""

    def test_cli_with_explicit_config(self, sample_md, tmp_path):
        config_path = tmp_path / "config.yaml"
        # Minimal config that works with FAKE_LLM seam
        config_path.write_text("""
project_id: "test"
source_lang: "en"
target_lang: "zh"
llm_pool:
  translation:
    - provider: "openai"
      model: "test-model"
      priority: 1
      role: "translation"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
    - provider: "openai"
      model: "test-model-2"
      priority: 2
      role: "translation"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
  judging:
    - provider: "openai"
      model: "test-model"
      priority: 1
      role: "judging"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
    - provider: "openai"
      model: "test-model-2"
      priority: 2
      role: "judging"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
  restoration:
    - provider: "openai"
      model: "test-model"
      priority: 1
      role: "restoration"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
    - provider: "openai"
      model: "test-model-2"
      priority: 2
      role: "restoration"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
  profiling:
    - provider: "openai"
      model: "test-model"
      priority: 1
      role: "profiling"
      api_key: "${ZHIPU_API_KEY}"
      base_url: "http://localhost:8080/v1"
""", encoding="utf-8")
        result = runner.invoke(app, [
            "profile-doc",
            str(sample_md),
            "--config", str(config_path),
        ])
        assert result.exit_code == 0


class TestProfileDocBinaryFormats:
    """T3.1 tests for profile-doc CLI with binary document formats."""

    def test_docx_file_is_supported_via_cli(self, tmp_path):
        """ol profile-doc should not crash on .docx files (T3.1 SURFACE check)."""
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx not installed")

        docx_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Hello from docx")
        doc.add_paragraph("This is a test paragraph for profile-doc.")
        doc.save(str(docx_path))

        result = runner.invoke(app, ["profile-doc", str(docx_path)])
        assert result.exit_code == 0, f"Exit: {result.exit_code}, output: {result.output}"
        # The output should be valid JSON with the StyleGuide fields
        try:
            output = json.loads(result.stdout)
        except json.JSONDecodeError:
            # If the output isn't JSON (e.g., error message), re-check
            pytest.fail(f"Output is not JSON: {result.stdout!r}")
        assert "tone" in output, f"Missing 'tone' in output: {output}"
        assert "summary" in output, f"Missing 'summary' in output: {output}"

    def test_pptx_file_is_supported_via_cli(self, tmp_path):
        """ol profile-doc should not crash on .pptx files."""
        try:
            from pptx import Presentation
        except ImportError:
            pytest.skip("python-pptx not installed")

        pptx_path = tmp_path / "test.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Hello from pptx"
        prs.save(str(pptx_path))

        result = runner.invoke(app, ["profile-doc", str(pptx_path)])
        assert result.exit_code == 0, f"Exit: {result.exit_code}, output: {result.output}"
        try:
            output = json.loads(result.stdout)
        except json.JSONDecodeError:
            pytest.fail(f"Output is not JSON: {result.stdout!r}")
        assert "tone" in output, f"Missing 'tone' in output: {output}"

    def test_real_docx_fixture_works_via_cli(self):
        """The real E2E test fixture (.docx) can be profiled via the CLI."""
        fixture = Path("爱上海尔_第二章_全球创牌 - E2E测试专用.docx")
        if not fixture.exists():
            pytest.skip(f"E2E fixture not found at {fixture}")

        result = runner.invoke(app, ["profile-doc", str(fixture)])
        assert result.exit_code == 0, f"Exit: {result.exit_code}, output: {result.output}"
        try:
            output = json.loads(result.stdout)
        except json.JSONDecodeError:
            pytest.fail(f"Output is not JSON: {result.stdout!r}")
        assert "tone" in output
        assert "register" in output
        assert "summary" in output
        # The summary should be non-empty (the docx has content)
        assert len(output.get("summary", "")) > 0, (
            f"Summary should be non-empty for a real .docx: {output.get('summary')!r}"
        )

    def test_unsupported_binary_format_gives_clear_error(self, tmp_path):
        """ol profile-doc on a fake .pdf gives a clear error message."""
        fake_pdf = tmp_path / "test.pdf"
        # Use bytes that are invalid UTF-8 to trigger UnicodeDecodeError
        fake_pdf.write_bytes(b"\xff\xfe\x80\x50\x44\x46")

        result = runner.invoke(app, ["profile-doc", str(fake_pdf)])
        # Exit code should be non-zero (error)
        assert result.exit_code != 0, f"Expected error exit code, got {result.exit_code}"
        # The error message should mention unsupported format
        combined = (result.output + (str(result.exception) if result.exception else "")).lower()
        assert "unsupported" in combined or "format" in combined or "binary" in combined or "cannot read" in combined, (
            f"Expected clear error about unsupported format, got: {combined[:300]}"
        )
