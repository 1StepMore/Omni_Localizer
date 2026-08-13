"""ol profile-doc — Profile a document's writing style.

Reads a document, calls the LLM-based doc_profiler, and prints the
resulting StyleGuide as JSON. With --output, writes to a file.
"""
from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Optional

import typer

from cli._shared import ExitCode, validate_input_file, warn_fake_llm_mode
from ol_style.cache import ProfileCache
from ol_style.doc_profiler import profile_document


def _extract_text_from_file(path: Path) -> str:
    """Extract text from a document file, handling binary formats.

    Supports: .md, .txt, .csv, .json, .xml, .html, .xlf, .yaml (UTF-8)
              .docx (via python-docx, optional)
              .pptx (via python-pptx, optional)

    Args:
        path: Path to the document file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If the file format is unsupported or extraction fails.
    """
    suffix = path.suffix.lower()

    if suffix in (".md", ".txt", ".csv", ".json", ".xml", ".html", ".htm",
                  ".xlf", ".xliff", ".yaml", ".yml"):
        return path.read_text(encoding="utf-8")

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError:
            raise ValueError(
                "python-docx is required to read .docx files. "
                "Install with: pip install python-docx"
            )
        doc = Document(str(path))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    if suffix == ".pptx":
        try:
            from pptx import Presentation
        except ImportError:
            raise ValueError(
                "python-pptx is required to read .pptx files. "
                "Install with: pip install python-pptx"
            )
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n\n".join(t for t in texts if t.strip())

    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        raise ValueError(
            f"Unsupported binary format: {suffix}. "
            f"Supported: .md, .txt, .docx, .pptx, .csv, .json, .xml, .html"
        )


def profile_doc(
    input_file: str = typer.Argument(
        ..., help="Path to input document (.md, .txt, .docx, .pptx, etc.)",
    ),
    output: Optional[str] = typer.Option(
        None, "--output", "-o",
        help="Path to write the JSON profile. If omitted, prints to stdout.",
    ),
    source_lang: str = typer.Option(
        "en", "--source-lang", "-s",
        help="Source language code (default: en)",
    ),
    config: Optional[str] = typer.Option(
        None, "--config", "-c",
        help="Path to OL YAML config (default: config/default.yaml)",
    ),
    cache_dir: Optional[str] = typer.Option(
        None, "--cache-dir",
        help="Directory to store profile cache. If omitted, in-memory only.",
    ),
) -> None:
    """Profile a document's writing style and emit a StyleGuide."""
    try:
        path = validate_input_file(input_file)
    except typer.BadParameter as e:
        typer.echo(f"Error: {e.message}", err=True)
        raise typer.Exit(code=ExitCode.CLI_USAGE_ERROR)

    warn_fake_llm_mode()

    try:
        content = _extract_text_from_file(path)
    except (OSError, ValueError) as e:
        typer.echo(f"Error reading input file: {e}", err=True)
        raise typer.Exit(code=ExitCode.PIPELINE_ERROR)

    cache: ProfileCache | None = None
    if cache_dir:
        cache = ProfileCache(cache_dir=Path(cache_dir))
    else:
        cache = ProfileCache()

    try:
        guide = asyncio.run(profile_document(
            content=content,
            source_lang=source_lang,
            config_path=config,
            cache=cache,
        ))
    except Exception as e:
        typer.echo(f"Error: profiling failed: {e}", err=True)
        raise typer.Exit(code=ExitCode.PIPELINE_ERROR)

    output_dict = guide.to_dict()
    output_json = json.dumps(output_dict, ensure_ascii=False, indent=2)
    if output:
        try:
            Path(output).write_text(output_json, encoding="utf-8")
        except OSError as e:
            typer.echo(f"Error writing output file: {e}", err=True)
            raise typer.Exit(code=ExitCode.PIPELINE_ERROR)
        typer.echo(f"Profile written to: {output}")
    else:
        typer.echo(output_json)
    raise typer.Exit(code=ExitCode.SUCCESS)
